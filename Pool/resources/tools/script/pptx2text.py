"""S-GAR tool — PPTX text extraction via python-pptx (https://python-pptx.readthedocs.io, MIT)."""
import sys
import os
from _sgar_cli import emit
from pptx import Presentation

if len(sys.argv) < 2:
    print("Usage: python pptx2text.py <file_path>", file=sys.stderr)
    sys.exit(1)
path = sys.argv[1]
if not os.path.exists(path):
    emit({"status": "error", "message": f"{path} not found"})
    sys.exit(2)
try:
    prs = Presentation(path)
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        texts = [sh.text for sh in slide.shapes if sh.has_text_frame and sh.text.strip()]
        slides.append({"slide": idx, "text": "\n".join(texts)})
    emit({"status": "success", "tool": "python-pptx", "target": path,
          "slide_count": len(slides), "slides": slides})
except Exception as exc:
    emit({"status": "error", "tool": "python-pptx", "message": f"{type(exc).__name__}: {exc}"})
    sys.exit(1)
