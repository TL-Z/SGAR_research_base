"""S-GAR tool — DOCX text extraction via python-docx (https://python-docx.readthedocs.io, MIT).
Extracts paragraph and table text from a real .docx file using the recognized
python-docx library."""
import sys
import os
from _sgar_cli import emit
import docx  # python-docx

if len(sys.argv) < 2:
    print("Usage: python docx2text.py <file_path>", file=sys.stderr)
    sys.exit(1)
path = sys.argv[1]
if not os.path.exists(path):
    emit({"status": "error", "message": f"{path} not found"})
    sys.exit(2)
try:
    document = docx.Document(path)
    paras = [p.text for p in document.paragraphs if p.text.strip()]
    tables = [
        [[cell.text for cell in row.cells] for row in tbl.rows]
        for tbl in document.tables
    ]
    emit({"status": "success", "tool": "python-docx", "target": path,
          "paragraph_count": len(paras), "text": "\n".join(paras), "tables": tables})
except Exception as exc:
    emit({"status": "error", "tool": "python-docx", "message": f"{type(exc).__name__}: {exc}"})
    sys.exit(1)
