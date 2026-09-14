#!/usr/bin/env python3
"""Run one real, isolated RC1 smoke case for every catalog-active Tool."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TOOLS = PROJECT_ROOT / "Pool" / "resources" / "json" / "tools.json"
DEFAULT_RUNTIME_LOCK = PROJECT_ROOT / "sgar_mvp" / "config" / "rc1_runtime_lock.json"
DEFAULT_OUTPUT_ROOT = PROJECT_ROOT / "sgar_mvp" / "execution_outputs" / "tool_smoke"

REST_ARGS: dict[str, list[str]] = {
    "tool.api.agify.v1": ["alex"],
    "tool.api.arxiv.search.v1": ["agent routing", "1"],
    "tool.api.clinicaltrials.v1": ["diabetes"],
    "tool.api.coingecko.price.v1": ["bitcoin", "usd"],
    "tool.api.crossref.works.v1": ["agent routing"],
    "tool.api.datamuse_words.v1": ["ocean"],
    "tool.api.dictionary.v1": ["example"],
    "tool.api.dog_ceo.v1": ["hound"],
    "tool.api.exchangerate.convert.v1": ["USD", "EUR", "1"],
    "tool.api.fda_drug.v1": ["aspirin"],
    "tool.api.frankfurter.rates.v1": ["USD", "EUR"],
    "tool.api.fruityvice.v1": ["banana"],
    "tool.api.genderize.v1": ["alex"],
    "tool.api.google_dns.v1": ["example.com", "A"],
    "tool.api.hackernews_item.v1": ["8863"],
    "tool.api.ip_api_geo.v1": ["8.8.8.8"],
    "tool.api.iss_location.v1": [],
    "tool.api.met_museum_art.v1": ["sunflowers"],
    "tool.api.nager_holidays.v1": ["2026", "US"],
    "tool.api.nationalize.v1": ["alex"],
    "tool.api.nominatim.geocode.v1": ["Shanghai Jiao Tong University"],
    "tool.api.openfoodfacts.product.v1": ["3017620422003"],
    "tool.api.opensky_flights.v1": [],
    "tool.api.opentrivia.v1": ["1"],
    "tool.api.open_meteo.v1": ["31.2304", "121.4737"],
    "tool.api.poetrydb.v1": ["Shakespeare"],
    "tool.api.pokeapi.v1": ["pikachu"],
    "tool.api.pubmed.esearch.v1": ["agent routing"],
    "tool.api.randomuser.v1": ["1"],
    "tool.api.rest_countries.info.v1": ["China"],
    "tool.api.rickandmorty.v1": ["character", "1"],
    "tool.api.stackexchange.search.v1": ["python", "stackoverflow"],
    "tool.api.thecocktaildb.v1": ["margarita"],
    "tool.api.themealdb.v1": ["Arrabiata"],
    "tool.api.thesportsdb.v1": ["Arsenal"],
    "tool.api.universities_hipolabs.v1": ["China"],
    "tool.api.usgs.earthquakes.v1": [],
    "tool.api.wikipedia.summary.v1": ["Artificial intelligence", "en"],
    "tool.api.worldbank_indicator.v1": ["CHN", "SP.POP.TOTL"],
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256(data: bytes) -> str:
    return "sha256:" + hashlib.sha256(data).hexdigest()


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def run(command: list[str], *, cwd: Path | None = None, timeout: int = 300) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        command,
        cwd=str(cwd) if cwd else None,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=timeout,
        check=False,
    )


def container_path(path: Path) -> str:
    return "/app/" + path.resolve().relative_to(PROJECT_ROOT).as_posix()


def build_fixture_template(root: Path, image: str) -> None:
    root.mkdir(parents=True, exist_ok=False)
    write(root / "sample.py", "import os\n\ndef add(a, b):\n    unused = 1\n    return a+b\n")
    write(root / "runner_target.py", "print('RC1_TOOL_SMOKE_OK')\n")
    write(root / "test_sample.py", "def test_add():\n    assert 1 + 1 == 2\n")
    write(root / "sample.md", "# RC1\n\n| name | value |\n| --- | --- |\n| alpha | 1 |\n")
    write(root / "sample.html", "<html><body><p>Hello RC1</p></body></html>\n")
    write(root / "sample.csv", "name,group,value\nalpha,A,1\nbeta,B,2\n")
    write(root / "sample.json", '{"name":"rc1","items":[1,2]}\n')
    write(root / "schema.json", '{"type":"object","required":["name"],"properties":{"name":{"type":"string"}}}\n')
    write(root / "sample.yaml", "name: rc1\nitems:\n  - 1\n  - 2\n")
    write(root / "sample.xml", "<root><item id=\"1\">RC1</item></root>\n")
    write(root / "sample.env", "RC1_NAME=test\nRC1_COUNT=1\n")
    write(root / "sample.graphql", "type Query { hello: String! }\n")
    write(root / "sample.sql", "CREATE TABLE users (id INTEGER PRIMARY KEY, name TEXT NOT NULL);\n")
    write(root / "sample.sh", "#!/bin/sh\necho RC1\n")
    write(root / "sample.c", "int main(){return 0;}\n")
    write(root / "sample.rs", "fn main(){println!(\"RC1\");}\n")
    write(root / "sample.go", "package main\nimport \"fmt\"\nfunc main(){fmt.Println(\"RC1\")}\n")
    write(root / "sample.js", "const value = 1;\nconsole.log(value);\n")
    write(root / "sample.css", "body { color: #000; }\n")
    write(root / "commit_message.txt", "test: verify rc1 tool smoke\n")
    write(root / "requirements.txt", "requests==2.32.3\n")
    write(
        root / "junit.xml",
        '<?xml version="1.0"?><testsuite tests="1" failures="0"><testcase classname="rc1" name="ok"/></testsuite>\n',
    )
    write(root / "edges.json", '[["a","b"],["b","c"]]\n')

    go = root / "go_project"
    write(go / "go.mod", "module example.com/rc1\n\ngo 1.22\n")
    write(go / "main.go", "package rc1\nfunc Add(a,b int) int { return a+b }\n")
    write(go / "main_test.go", "package rc1\nimport \"testing\"\nfunc TestAdd(t *testing.T){if Add(1,1)!=2{t.Fail()}}\n")
    rust = root / "rust_project"
    write(rust / "Cargo.toml", '[package]\nname="rc1_smoke"\nversion="0.1.0"\nedition="2021"\n')
    write(rust / "src" / "lib.rs", "pub fn add(a:i32,b:i32)->i32{a+b}\n#[cfg(test)] mod tests {use super::*; #[test] fn ok(){assert_eq!(add(1,1),2);}}\n")
    js = root / "js_project"
    write(
        js / "package.json",
        json.dumps(
            {
                "name": "rc1-smoke",
                "version": "1.0.0",
                "private": True,
                "scripts": {"test": "node -e \"if (1+1!==2) process.exit(1)\""},
            },
            indent=2,
        )
        + "\n",
    )
    write(js / "package-lock.json", '{"name":"rc1-smoke","version":"1.0.0","lockfileVersion":3,"packages":{"":{"name":"rc1-smoke","version":"1.0.0"}}}\n')
    write(js / "sum.test.js", "test('sum', () => { expect(1 + 1).toBe(2); });\n")
    write(js / "sum.test.mjs", "import { test, expect } from 'vitest';\ntest('sum',()=>expect(1+1).toBe(2));\n")
    write(js / "index.ts", "const value: number = 1; console.log(value);\n")
    write(js / "tsconfig.json", '{"compilerOptions":{"strict":true,"noEmit":true},"include":["index.ts"]}\n')

    repo = root / "repo"
    repo.mkdir()
    git_commands = (
        ["git", "init"],
        ["git", "config", "user.email", "rc1@example.invalid"],
        ["git", "config", "user.name", "SGAR RC1"],
    )
    for command in git_commands:
        proc = run(command, cwd=repo)
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr or proc.stdout)
    write(repo / "tracked.py", "print('tracked')\n")
    if run(["git", "add", "tracked.py"], cwd=repo).returncode != 0:
        raise RuntimeError("Unable to stage Git smoke fixture")
    proc = run(["git", "commit", "-m", "test: initialize rc1 fixture"], cwd=repo)
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr or proc.stdout)
    write(repo / "untracked.txt", "untracked\n")

    binary_code = r'''
from pathlib import Path
import sys
root=Path(sys.argv[1])
from docx import Document
d=Document(); d.add_paragraph("RC1 document"); d.save(root/"sample.docx")
from pptx import Presentation
p=Presentation(); s=p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text="RC1"; s.placeholders[1].text="Tool smoke"; p.save(root/"sample.pptx")
from openpyxl import Workbook
w=Workbook(); ws=w.active; ws.append(["name","value"]); ws.append(["alpha",1]); w.save(root/"sample.xlsx")
from PIL import Image
Image.new("RGB",(2,2),(255,0,0)).save(root/"sample.png")
from pypdf import PdfWriter
writer=PdfWriter(); writer.add_blank_page(width=72,height=72)
with (root/"sample.pdf").open("wb") as f: writer.write(f)
'''
    proc = run(
        [
            "docker",
            "run",
            "--rm",
            "-v",
            f"{PROJECT_ROOT}:/app",
            image,
            "python",
            "-c",
            binary_code,
            container_path(root),
        ],
        timeout=300,
    )
    if proc.returncode != 0:
        raise RuntimeError(f"Binary fixture generation failed: {proc.stderr or proc.stdout}")


def library_args(rid: str, root: Path) -> list[str] | None:
    p = lambda name: container_path(root / name)
    mapping = {
        "tool.lib.beautifulsoup_scrape.v1": [p("sample.html"), "p"],
        "tool.lib.bleach_sanitize.v1": ["<b>safe</b><script>bad()</script>"],
        "tool.lib.chardet_detect.v1": [p("sample.md")],
        "tool.lib.cryptography_hash_sign.v1": ["rc1", "hash"],
        "tool.lib.dateparser_parse.v1": ["2026-08-06"],
        "tool.lib.detect_secrets.v1": [p("sample.py")],
        "tool.lib.faker_generate.v1": ["basic", "1"],
        "tool.lib.genson_schema.v1": [p("sample.json")],
        "tool.lib.jinja2_render.v1": ["Hello {{ name }}", '{"name":"RC1"}'],
        "tool.lib.jsonschema_validate.v1": [p("sample.json"), p("schema.json")],
        "tool.lib.langdetect_detect.v1": ["This is a sufficiently long English sentence for detection."],
        "tool.lib.markdown_it_render.v1": ["# RC1"],
        "tool.lib.networkx_analyze.v1": ['[["a","b"],["b","c"]]', "degree"],
        "tool.lib.numpy_linalg.v1": ["[[1,0],[0,1]]", "det"],
        "tool.lib.pandas_describe.v1": [p("sample.csv")],
        "tool.lib.pandas_pivot.v1": [p("sample.csv"), "group", "value"],
        "tool.lib.pdfplumber_extract.v1": [p("sample.pdf")],
        "tool.lib.phonenumbers_parse.v1": ["+8613800138000", "CN"],
        "tool.lib.pillow_transform.v1": [p("sample.png"), "grayscale"],
        "tool.lib.pint_convert.v1": ["1", "meter", "centimeter"],
        "tool.lib.pygments_highlight.v1": ["print('RC1')", "python"],
        "tool.lib.pygount_sloc.v1": [p("sample.py")],
        "tool.lib.pypdf_merge.v1": [p("sample.pdf")],
        "tool.lib.qrcode_generate.v1": ["RC1"],
        "tool.lib.radon_complexity.v1": [p("sample.py")],
        "tool.lib.shapely_geometry.v1": ["POINT (0 0)", "POINT (1 1)", "distance"],
        "tool.lib.sqlglot_transpile.v1": ["SELECT 1", "sqlite", "postgres"],
        "tool.lib.sympy_solve.v1": ["x**2-1"],
        "tool.lib.tablib_convert.v1": [p("sample.csv"), "json"],
        "tool.lib.validators_check.v1": ["https://example.com", "url"],
    }
    return mapping.get(rid)


def script_args(rid: str, root: Path) -> list[str] | None:
    p = lambda name: container_path(root / name)
    mapping = {
        "tool.bandit_security_scanner.v1": [p("sample.py")],
        "tool.black_format_runner.v1": [p("sample.py")],
        "tool.cargo_test_runner.v1": [p("rust_project")],
        "tool.clang_format_checker.v1": [p("sample.c")],
        "tool.cors_headers_checker.v1": ["https://example.com"],
        "tool.csv_column_filter.v1": [p("sample.csv"), "name"],
        "tool.dependency_cve_checker.v1": [p("requirements.txt")],
        "tool.dns_query_latency_checker.v1": ["example.com", "A"],
        "tool.docx2text.v1": [p("sample.docx")],
        "tool.dotenv_to_json_converter.v1": [p("sample.env")],
        "tool.eslint_checker.v1": [p("sample.js")],
        "tool.eslint_code_formatter.v1": [p("sample.js")],
        "tool.flake8_lint_runner.v1": [p("sample.py")],
        "tool.git_blame_analyzer.v1": [p("repo/tracked.py")],
        "tool.git_branch_cleaner.v1": [p("repo")],
        "tool.git_commit_lint.v1": [p("commit_message.txt")],
        "tool.git_diff_stat.v1": [p("repo")],
        "tool.git_log_json.v1": [p("repo"), "1"],
        "tool.git_log_pretty_format.v1": [p("repo"), "1"],
        "tool.git_merge_conflicts_scanner.v1": [p("repo")],
        "tool.git_submodule_status.v1": [p("repo")],
        "tool.git_tags_inspector.v1": [p("repo")],
        "tool.go_fmt_checker.v1": [p("sample.go")],
        "tool.go_test_coverage.v1": [p("go_project")],
        "tool.go_test_runner.v1": [p("go_project")],
        "tool.graphql_schema_validator.v1": [p("sample.graphql")],
        "tool.html_to_markdown_converter.v1": [p("sample.html")],
        "tool.http_headers_security_auditor.v1": ["https://example.com"],
        "tool.jest_test_runner.v1": [p("js_project")],
        "tool.json_jq_query_executor.v1": [p("sample.json"), ".items"],
        "tool.json_to_toml_converter.v1": [p("sample.json")],
        "tool.json_to_yaml.v1": [p("sample.json")],
        "tool.junit_xml_parser.v1": [p("junit.xml")],
        "tool.markdown_table_extractor.v1": [p("sample.md")],
        "tool.markdown_to_html_converter.v1": [p("sample.md")],
        "tool.memory_profiler_exporter.v1": [p("runner_target.py")],
        "tool.npm_audit_json_exporter.v1": [p("js_project/package.json")],
        "tool.npm_package_auditor.v1": [p("js_project/package.json")],
        "tool.npm_test_runner.v1": [p("js_project")],
        "tool.pptx2text.v1": [p("sample.pptx")],
        "tool.pytest_runner.v1": [p("test_sample.py")],
        "tool.python_dead_code_scanner.v1": [p("sample.py")],
        "tool.python_script_runner.v1": [p("runner_target.py")],
        "tool.python_static_analyzer.v1": [p("sample.py")],
        "tool.ruff_lint_runner.v1": [p("sample.py")],
        "tool.rustfmt_checker.v1": [p("sample.rs")],
        "tool.shellcheck_runner.v1": [p("sample.sh")],
        "tool.sql_ddl_to_json_schema.v1": [p("sample.sql")],
        "tool.stylelint_css_checker.v1": [p("sample.css")],
        "tool.system_ram_free_checker.v1": [],
        "tool.tsc_compile_runner.v1": [p("js_project")],
        "tool.vitest_runner.v1": [p("js_project")],
        "tool.xlsx2csv.v1": [p("sample.xlsx")],
        "tool.xml_to_json.v1": [p("sample.xml")],
        "tool.xml_xpath_query_executor.v1": [p("sample.xml"), "//item"],
        "tool.yaml_to_json.v1": [p("sample.yaml")],
    }
    return mapping.get(rid)


def mcp_args(rid: str, root: Path) -> list[str] | None:
    p = lambda name: container_path(root / name)
    repo = p("repo")
    mapping = {
        "tool.mcp.fs_create_directory.v1": [p("created")],
        "tool.mcp.fs_directory_tree.v1": [p(".")],
        "tool.mcp.fs_edit_file.v1": [p("sample.md"), "RC1", "RC1_EDITED"],
        "tool.mcp.fs_get_file_info.v1": [p("sample.md")],
        "tool.mcp.fs_list_directory.v1": [p(".")],
        "tool.mcp.fs_move_file.v1": [p("sample.env"), p("moved.env")],
        "tool.mcp.fs_read_file.v1": [p("sample.md")],
        "tool.mcp.fs_read_multiple.v1": [p("sample.md") + "," + p("sample.json")],
        "tool.mcp.fs_search_files.v1": [p("."), "*.md"],
        "tool.mcp.fs_write_file.v1": [p("written.txt"), "RC1"],
        "tool.mcp.git_add.v1": [repo, "untracked.txt"],
        "tool.mcp.git_branch_list.v1": [repo],
        "tool.mcp.git_commit.v1": [repo, "test: rc1 smoke commit"],
        "tool.mcp.git_create_branch.v1": [repo, "rc1-smoke-branch"],
        "tool.mcp.git_diff_staged.v1": [repo],
        "tool.mcp.git_diff_unstaged.v1": [repo],
        "tool.mcp.git_log.v1": [repo, "1"],
        "tool.mcp.git_show.v1": [repo, "HEAD"],
        "tool.mcp.git_status.v1": [repo],
        "tool.mcp.sequential_thinking.v1": ["Verify one RC1 smoke step."],
    }
    return mapping.get(rid)


def prepare_tool_workspace(rid: str, template: Path, workspace: Path) -> None:
    shutil.copytree(template, workspace)
    repo = workspace / "repo"
    if rid in {"tool.mcp.git_commit.v1", "tool.mcp.git_diff_staged.v1"}:
        proc = run(["git", "add", "untracked.txt"], cwd=repo)
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr or proc.stdout)
    if rid == "tool.mcp.git_diff_unstaged.v1":
        write(repo / "tracked.py", "print('changed')\n")


def parse_output(stdout: str) -> dict[str, Any] | None:
    text = (stdout or "").strip()
    candidates = [text]
    if "\n" in text:
        candidates.extend(reversed(text.splitlines()))
    for candidate in candidates:
        try:
            payload = json.loads(candidate)
        except Exception:
            continue
        if isinstance(payload, dict):
            return payload
    return None


def classify_result(
    *,
    return_code: int,
    stdout: str,
    stderr: str,
    network_required: bool,
) -> tuple[str, bool, str]:
    payload = parse_output(stdout)
    combined = f"{stdout}\n{stderr}".lower()
    if return_code == 0 and payload:
        status = str(payload.get("status") or payload.get("result") or "").lower()
        nested = payload.get("result") if isinstance(payload.get("result"), dict) else {}
        if nested.get("is_error") is True:
            return "blocked", False, "mcp_tool_reported_error"
        if status in {"success", "ok", "passed", "valid", "succeeded"}:
            http_status = payload.get("http_status")
            if http_status is not None and not 200 <= int(http_status) < 400:
                return "blocked", False, f"http_status_{http_status}"
            return "ready", True, ""
    transient_markers = (
        "timeout",
        "timed out",
        "temporary failure",
        "connection reset",
        "connection refused",
        "name resolution",
        "sslerror",
        "unexpected_eof",
        "max retries exceeded",
        "429",
        "502",
        "503",
        "504",
    )
    if network_required and any(marker in combined for marker in transient_markers):
        return "transient_failure", False, "network_or_provider_transient"
    return "blocked", False, "real_smoke_failed"


def smoke_args(manifest: dict[str, Any], workspace: Path) -> list[str] | None:
    runtime = str((manifest.get("execution") or {}).get("runtime") or "")
    rid = str(manifest["resource_id"])
    if runtime == "rest_api":
        return REST_ARGS.get(rid)
    if runtime == "python_library":
        return library_args(rid, workspace)
    if runtime == "python_script":
        return script_args(rid, workspace)
    if runtime == "mcp_server":
        return mcp_args(rid, workspace)
    return None


def run_one(
    manifest: dict[str, Any],
    *,
    image: str,
    image_id: str,
    template: Path,
    workspaces: Path,
    evidence_dir: Path,
    timeout_seconds: int,
    transient_retries: int,
) -> dict[str, Any]:
    rid = str(manifest["resource_id"])
    workspace = workspaces / rid.replace(".", "_")
    prepare_tool_workspace(rid, template, workspace)
    arguments = smoke_args(manifest, workspace)
    if arguments is None:
        return {
            "resource_id": rid,
            "status": "blocked",
            "smoke_passed": False,
            "failure_type": "smoke_case_missing",
            "attempt_count": 0,
            "checked_at": utc_now(),
        }
    uri = str((manifest.get("execution") or {}).get("uri") or "")
    wrapper = (PROJECT_ROOT / uri.removeprefix("file://")).resolve()
    declared_hash = str((manifest.get("provenance") or {}).get("source_hash") or "")
    actual_hash = sha256(wrapper.read_bytes()) if wrapper.is_file() else ""
    if declared_hash != actual_hash:
        return {
            "resource_id": rid,
            "status": "blocked",
            "smoke_passed": False,
            "failure_type": "source_hash_mismatch",
            "attempt_count": 0,
            "checked_at": utc_now(),
        }
    requirements = manifest.get("runtime_requirements") or {}
    if requirements.get("docker_image") != image or requirements.get("install_policy") != "never":
        return {
            "resource_id": rid,
            "status": "blocked",
            "smoke_passed": False,
            "failure_type": "runtime_policy_mismatch",
            "attempt_count": 0,
            "checked_at": utc_now(),
        }

    max_attempts = max(1, transient_retries + 1)
    final: dict[str, Any] = {}
    for attempt in range(1, max_attempts + 1):
        started = time.perf_counter()
        command = [
            "docker",
            "run",
            "--rm",
            "-e",
            f"SGAR_WORKSPACE_ROOT={container_path(workspace)}",
            "-e",
            "SGAR_HOST_WORKSPACE_ROOT=/app",
            "-v",
            f"{PROJECT_ROOT}:/app",
            "-w",
            "/app",
            image,
            "python",
            container_path(wrapper),
            *arguments,
        ]
        try:
            proc = run(command, timeout=timeout_seconds)
            stdout = proc.stdout or ""
            stderr = proc.stderr or ""
            status, passed, failure_type = classify_result(
                return_code=proc.returncode,
                stdout=stdout,
                stderr=stderr,
                network_required=bool(requirements.get("network_required")),
            )
            evidence_dir.mkdir(parents=True, exist_ok=True)
            evidence_stem = rid.replace(".", "_")
            stdout_path = evidence_dir / f"{evidence_stem}.stdout.txt"
            stderr_path = evidence_dir / f"{evidence_stem}.stderr.txt"
            stdout_path.write_text(stdout, encoding="utf-8")
            stderr_path.write_text(stderr, encoding="utf-8")
            final = {
                "resource_id": rid,
                "status": status,
                "smoke_passed": passed,
                "failure_type": failure_type or None,
                "return_code": proc.returncode,
                "latency_ms": round((time.perf_counter() - started) * 1000, 2),
                "attempt_count": attempt,
                "checked_at": utc_now(),
                "source_hash": actual_hash,
                "output_sha256": sha256(stdout.encode("utf-8", errors="replace")),
                "output_bytes": len(stdout.encode("utf-8", errors="replace")),
                "stderr_preview": stderr[:500] or None,
                "stdout_artifact": str(stdout_path.relative_to(PROJECT_ROOT).as_posix()),
                "stderr_artifact": str(stderr_path.relative_to(PROJECT_ROOT).as_posix()),
                "docker_image_id": image_id,
            }
        except subprocess.TimeoutExpired:
            status = "transient_failure" if requirements.get("network_required") else "blocked"
            final = {
                "resource_id": rid,
                "status": status,
                "smoke_passed": False,
                "failure_type": "smoke_timeout",
                "latency_ms": timeout_seconds * 1000,
                "attempt_count": attempt,
                "checked_at": utc_now(),
                "source_hash": actual_hash,
                "docker_image_id": image_id,
            }
        if final["status"] != "transient_failure" or attempt >= max_attempts:
            break
        time.sleep(1)
    return final


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--tools", type=Path, default=DEFAULT_TOOLS)
    parser.add_argument("--runtime-lock", type=Path, default=DEFAULT_RUNTIME_LOCK)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument("--tool-id", action="append")
    parser.add_argument("--timeout-sec", type=int, default=360)
    parser.add_argument("--transient-retries", type=int, default=1)
    parser.add_argument("--keep-workspaces", action="store_true")
    parser.add_argument("--fail-on-nonready", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    args.output_root = args.output_root.resolve()
    tools = json.loads(args.tools.read_text(encoding="utf-8-sig"))
    runtime_lock = json.loads(args.runtime_lock.read_text(encoding="utf-8-sig"))
    image = str(runtime_lock.get("image") or "")
    image_id = str(runtime_lock.get("image_id") or "")
    if image != "sgar-runtime:rc1" or not image_id:
        raise SystemExit("A sealed sgar-runtime:rc1 lock is required before Tool smoke tests")
    inspect = run(["docker", "image", "inspect", image])
    if inspect.returncode != 0:
        raise SystemExit(inspect.stderr or inspect.stdout)
    actual_image_id = str(json.loads(inspect.stdout)[0].get("Id") or "")
    if actual_image_id != image_id:
        raise SystemExit("Docker image ID differs from the sealed RC1 runtime lock")

    requested = set(args.tool_id or [])
    active = [
        item
        for item in tools
        if str(item.get("status") or "").lower() == "active"
        and str((item.get("execution") or {}).get("execution_status") or "active").lower()
        == "active"
        and (not requested or item["resource_id"] in requested)
    ]
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    run_dir = args.output_root / f"run_{timestamp}"
    template = run_dir / "fixture_template"
    workspaces = run_dir / "workspaces"
    evidence_dir = run_dir / "evidence"
    run_dir.mkdir(parents=True, exist_ok=False)
    workspaces.mkdir()
    build_fixture_template(template, image)

    missing_cases = [item["resource_id"] for item in active if smoke_args(item, template) is None]
    if missing_cases:
        raise RuntimeError(f"Active Tools without smoke cases: {missing_cases}")

    results: list[dict[str, Any]] = []
    for index, manifest in enumerate(active, start=1):
        result = run_one(
            manifest,
            image=image,
            image_id=image_id,
            template=template,
            workspaces=workspaces,
            evidence_dir=evidence_dir,
            timeout_seconds=args.timeout_sec,
            transient_retries=args.transient_retries,
        )
        results.append(result)
        print(
            f"[{index:03d}/{len(active):03d}] {result['status']:17} "
            f"{result['resource_id']}"
        )
    results.sort(key=lambda item: item["resource_id"])
    payload = {
        "schema_version": 1,
        "generated_at": utc_now(),
        "docker_image": image,
        "docker_image_id": image_id,
        "policy": {
            "real_execution_only": True,
            "dynamic_install": False,
            "max_auto_heals": 0,
            "filesystem_scope": "per-tool isolated workspace",
        },
        "summary": {
            "catalog_active": len(active),
            "ready": sum(item["status"] == "ready" for item in results),
            "blocked": sum(item["status"] == "blocked" for item in results),
            "transient_failure": sum(
                item["status"] == "transient_failure" for item in results
            ),
        },
        "tools": results,
    }
    report = run_dir / "report.json"
    report.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    latest = args.output_root / "tool_smoke_latest.json"
    latest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(report, latest)
    if not args.keep_workspaces:
        shutil.rmtree(workspaces, ignore_errors=True)
        shutil.rmtree(template, ignore_errors=True)
    print(json.dumps(payload["summary"], indent=2))
    print(f"Report: {report}")
    nonready = any(item["status"] != "ready" for item in results)
    return 2 if args.fail_on_nonready and nonready else 0


if __name__ == "__main__":
    raise SystemExit(main())
